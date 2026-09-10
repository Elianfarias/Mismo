from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\elian\Mismo')
SRC=Path(r'C:\Users\elian\OneDrive\Documentos\Nuevo Juego\Mismo_GDD_v0.4_consolidado.docx')
D=Document(SRC)
# Retain the complete existing prototype specification, replacing outdated progression.
body=D._element.body
cut=False
for el in list(body):
    if el.tag==qn('w:sectPr'): continue
    if ''.join(el.itertext()).startswith('8 Progresión pendiente'):
        cut=True
    if cut: body.remove(el)
for p in D.paragraphs:
    t=p.text
    if t.startswith('Versión 0.4'): p.text='Versión 0.5 consolidada · 8 de septiembre de 2026'
    elif t.startswith('Documento de referencia del diseño'):
        p.text='Documento de diseño de Mismo. Conserva la especificación del prototipo v0.4 e incorpora la progresión del personaje, maestría, tiers de armas, transmog, skins pagas, heridas y escalado regional. Las decisiones nuevas se describen como pendientes de implementación.'
    elif t.startswith('Esta es la versión final'):
        p.text='La v0.5 establece la dirección de progresión y personalización. El estado implementado corresponde al prototipo documentado en la v0.4; las fórmulas, límites y valores nuevos requieren balance y pruebas.'
    elif t.startswith('Progreso global del personaje y de sus armas;'):
        p.text='Progreso del personaje mediante atributos; maestría por familia de arma; drops con tiers y modificadores; materiales, builds, dungeons y secretos; nuevas regiones, biomas, criaturas y posibles monturas. Los pueblos podrán ofrecer comercio, herrería y almacenamiento. Las secciones 8 a 12 definen la dirección acordada y los pendientes.'
    elif t.startswith('El progreso futuro del arma'):
        pass
    elif t.startswith('La base utiliza definiciones'):
        pass
    elif t.startswith('Assets/Data/Weapons contiene'):
        p.text=t.replace('esta consolidación incorpora las correcciones recientes.', 'las secciones 1 a 7 conservan el estado documentado en la v0.4; las secciones posteriores establecen el diseño futuro de la v0.5.')
for t in D.tables:
    for row in t.rows:
        if row.cells[0].text=='Progresión y guardado':
            row.cells[1].text='Diseño ampliado en v0.5. Atributos, maestría, drops, transmog, heridas y guardado pendientes de implementación.'
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    D.styles[name].font.color.rgb=RGBColor(0,0,0)
for node in list(D._element.iter(qn('w:pBdr'))) + list(D.styles.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)

def p(text):
    previous=list(D._element.body)[-2]
    para=D.add_paragraph(text)
    if previous.tag==qn('w:tbl'): para.paragraph_format.space_before=Pt(7)
def h(text): D.add_heading(text,2)
def page(title):
    D.add_heading(title,1).paragraph_format.page_break_before=True
def table(headers, rows, widths):
    t=D.add_table(rows=1,cols=len(headers)); t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,values in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for c,value,w in zip(cells,values,widths):
            c.text=value; c.width=Inches(w); c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DFE8ED' if i==0 else ('F5F7F8' if i%2==0 else 'FFFFFF')); pr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(3); para.paragraph_format.space_after=Pt(3)
                for r in para.runs: r.font.size=Pt(10); r.bold=i==0
        rp=t.rows[i]._tr.get_or_add_trPr(); rp.append(OxmlElement('w:cantSplit'))
        if i==0: rp.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)

# Remove the old explicit page-break paragraph immediately before section 8.
while D.paragraphs and not D.paragraphs[-1].text.strip():
    el=D.paragraphs[-1]._element; el.getparent().remove(el)

page('8 Personaje maestría y roles')
p('Diseño acordado, pendiente de implementación. La progresión se divide en atributos generales del personaje, calidad del ejemplar y maestría de su familia. Cada capa debe conservar una función reconocible.')
table(['Capa','Decisión del jugador','Persistencia'],[
 ['Personaje','Elegir atributos al subir de nivel.','Pertenece al personaje.'],
 ['Ejemplar de arma','Elegir drops por tier y modificadores.','Pertenece al objeto.'],
 ['Maestría de familia','Elegir daño o velocidad de ataque.','Se conserva al cambiar de ejemplar.']],[1.35,2.8,2.5])
h('Atributos del personaje')
p('Cada nivel permite elegir aumentos de atributos, inicialmente vida máxima, ataque y defensa o armadura. El catálogo final, los puntos por nivel, las fórmulas y la redistribución están pendientes. Defensa y armadura deben unificarse como término si representan la misma estadística.')
h('Maestría de armas')
p('La maestría se aprende por familia: progresar con espada se conserva al encontrar otra espada y no otorga automáticamente dominio de arco. Al subir maestría se elige entre daño del arma y velocidad de ataque. Alcance queda excluido; por ahora no se añaden otros atributos.')
p('Los límites de daño y velocidad se balancean por separado. Falta definir qué fases de cada ataque modifica la velocidad, cómo afecta a la carga del arco y cómo se combina con otros modificadores. La preparación debe seguir siendo relevante.')
p('Se conserva la dirección de ofrecer temprano el kit representativo y reconocer la contribución efectiva del arma que actúa, incluidas defensas útiles. Llevarla en la espalda no basta para ganar maestría. Curva, máximo, atribución de efectos persistentes y experiencia defensiva siguen pendientes. Las variantes de habilidades por hitos de la v0.4 ya no son la progresión principal acordada.')
h('Rol de la combinación')
p('Las armas definen acciones, ritmo, alcance y oportunidades. La pareja elegida, los atributos y futuros efectos de especialización determinan el rol de la build, sin clases tradicionales. Gemas y variantes de habilidades continúan como posibilidades, no como sistemas cerrados.')

page('9 Tiers y modificadores de armas')
p('Diseño acordado, pendiente de implementación. Los drops tienen tiers que mejoran su presupuesto de poder. Los modificadores distribuyen ventajas y penalizaciones; un tier superior no obliga a mejorar todas las estadísticas a la vez.')
h('Atributos y alcance')
p('Un ejemplar puede modificar daño, velocidad de ataque, vida máxima y armadura. Los efectos sobre vida y armadura afectan al personaje. El alcance es una propiedad propia del diseño del arma: no se obtiene por maestría ni como modificador aleatorio del drop.')
table(['Ejemplo de variante','Bonificación','Penalización'],[
 ['Coloso','Daño y vida máxima.','Velocidad de ataque.'],
 ['Duelista','Velocidad de ataque.','Armadura.'],
 ['Guardián','Armadura.','Daño.']],[1.7,2.55,2.4])
p('Los nombres y combinaciones de la tabla son ejemplos de balance, no un catálogo definitivo. Una misma variante puede existir en distintos tiers.')
h('Dos armas equipadas')
p('Los modificadores de vida máxima y armadura de ambas armas permanecen activos mientras estén equipadas, incluida la secundaria. Daño y velocidad se aplican al arma correspondiente. Alternar no cambia la vida máxima ni permite evadir una penalización guardando el arma.')
p('Reemplazar ejemplares sigue limitado a fuera de combate. Debe definirse cómo se ajustan vida actual y heridas al cambiar la vida máxima, evitando que reequipar genere curación gratuita. Las reglas existentes de cooldown continúan al alternar y reequipar.')
h('Pendientes de balance')
p('Faltan cantidad y nombres de tiers, rangos de atributos, probabilidades de drop, requisitos de uso y reglas de acumulación. La mejora de ejemplares mediante materiales, contemplada en v0.4, no queda confirmada como una cuarta capa de poder.')
p('El balance debe evaluar en conjunto ataque del personaje, daño del ejemplar y maestría. La velocidad necesita revisar daño por segundo, exposición y frecuencia de efectos; su valor no debe compararse con daño solo por el porcentaje visible.')

page('10 Transmog y skins pagas')
p('Diseño acordado, pendiente de implementación. La personalización de armas incluye apariencias desbloqueadas mediante destrucción de ejemplares y skins compradas con dinero real. Ambas vías son cosméticas y conservan las estadísticas y el alcance real del arma equipada.')
h('Desbloquear una apariencia mediante un arma')
p('El jugador entrega y destruye un ejemplar para aprender permanentemente su apariencia. El arma consumida no produce materiales ni oro. La elección enfrenta tres destinos posibles del objeto: conservar su apariencia, convertirlo en materiales o venderlo, cuando esos servicios estén disponibles.')
p('La interfaz debe mostrar qué ejemplar se consumirá y qué apariencia se desbloqueará. La operación de consumo y desbloqueo debe guardarse junta para evitar pérdidas o duplicaciones. Falta definir el tratamiento de apariencias ya aprendidas y el alcance de la colección entre personajes.')
h('Skins compradas con dinero real')
p('El término pagos se refiere expresamente a skins pagas con dinero real, no a una tarifa de oro. La compra desbloquea una apariencia cosmética; no aumenta tier, atributos, maestría ni alcance. No se exige destruir un arma para obtener una skin comprada.')
p('Precios, catálogo, plataforma de venta, vinculación de compras y disponibilidad permanecen pendientes. La tienda y el sistema de compras no forman parte del prototipo existente.')
h('Aplicación y lectura del combate')
p('La propuesta inicial es aplicar apariencias dentro de la misma familia de arma. Falta confirmar esta restricción y la compatibilidad entre modelos. La apariencia debe conservar una lectura coherente de tamaño, alcance, agarre y ataques.')
p('No se ha acordado cobrar oro por aplicar una apariencia. El costo confirmado de la vía obtenida jugando es destruir el ejemplar; la vía comercial utiliza dinero real.')

page('11 Contraataque curación y heridas')
h('Contraataque después del parry')
p('Cambio solicitado, pendiente de implementación: después de aplicar un parry, el jugador debe poder atacar y cortar la animación defensiva. La regla propuesta para probarlo habilita la cancelación de recuperación tras un parry exitoso; un intento fallido conserva su recuperación.')
p('Se propone un pequeño buffer para aceptar la intención de ataque próxima a la confirmación del bloqueo. Falta fijar la ventana y las acciones permitidas. Esta excepción no habilita alternancia libre durante ataques ni modifica automáticamente las reglas actuales de cancelación del arco.')
h('Daño que reduce la vida recuperable')
p('Diseño acordado, pendiente de implementación: una fracción del daño recibido se acumula como herida y deja de ser recuperable mediante curación normal. El 10 % es un ejemplo de prueba. El objetivo es limitar la curación sostenida y dar peso al daño acumulado.')
p('Regla propuesta de cálculo: usar daño real descontado de vida, después de mitigación. Herida añadida = daño recibido × porcentaje de herida. Máximo curable = vida máxima efectiva − heridas acumuladas, con límites que eviten valores inválidos.')
table(['Ejemplo con 100 de vida máxima','Resultado'],[
 ['Daño recibido','30 puntos.'],['Vida actual tras el golpe','70 puntos.'],['Herida al 10 %','3 puntos.'],['Máximo recuperable','97 puntos.']],[3.9,2.75])
p('Descansar en un pueblo o punto seguro es la propuesta para recuperar heridas, todavía sin confirmar. También faltan reglas para muerte, cambios de vida máxima, pisos de vida recuperable y compatibilidad con el secreto que cura. El HUD deberá distinguir vida actual, vida curable y heridas.')
h('Gemas y armas de curación')
p('Se mantiene como posibilidad una gema que permita curar a otros al causar daño. No se confirma todavía una gema concreta ni un kit sanador. Activaciones ligadas a aperturas, un cooldown compartido y una variante útil en solitario son opciones para probar. Las heridas no sustituyen el balance de recursos, frecuencia y cantidad de curación.')

page('12 Niveles de monstruos por región')
p('Diseño acordado, pendiente de implementación. Al descubrir y cargar por primera vez una región nueva se fija su nivel mínimo de monstruos a partir del nivel del personaje de ese momento y un incremento configurable.')
h('Asignación inicial y persistencia')
p('Mínimo regional = nivel del personaje al descubrir la región + incremento regional.')
p('El incremento puede variar; 15 o 20 niveles son ejemplos. El mínimo se registra una sola vez para la región completa. Cargar nuevos chunks de esa misma región, volver a visitarla o subir de nivel no vuelve a sumarlo ni a calcularlo.')
h('Escalado posterior')
p('Nivel de los monstruos = máximo entre el mínimo regional guardado y el nivel actual del personaje. Cuando el personaje supera el mínimo, los monstruos lo acompañan. La referencia es el nivel de PERSONAJE; tier y maestría no determinan este escalado.')
table(['Situación con incremento de 15','Nivel de monstruos'],[
 ['Región descubierta con personaje nivel 10','25; el mínimo guardado es 25.'],
 ['Regreso con personaje nivel 18','25; el mínimo no cambia.'],
 ['Regreso con personaje nivel 30','30; el mínimo sigue siendo 25.']],[4.0,2.65])
h('Descubrimiento y límites pendientes')
p('Debe existir una identidad persistente por región y un evento único de descubrimiento. Falta definir si una precarga distante cuenta como descubrimiento o si se requiere proximidad o entrada. Un chunk adicional de una región conocida nunca crea otro mínimo.')
p('Faltan nivel de la región inicial, incrementos por región, conversión de niveles a estadísticas y momento de actualización de enemigos ya presentes. El comportamiento cooperativo también está pendiente, incluida la referencia de nivel al descubrir una región.')
p('El escalado mantiene relevantes regiones anteriores, pero puede reducir la sensación de superioridad al volver. Las pruebas deben verificar que progresar siga aportando valor y que el incremento inicial permita un desafío abordable. El streaming de nuevas regiones no está implementado en el paisaje actual guardado desde el editor.')

page('13 Implementación y decisiones pendientes')
p('La v0.5 es una actualización de diseño. No implica cambios en el código ni una nueva build. Se conserva el prototipo descrito en las secciones 1 a 7 y se ordena el trabajo futuro para validar las nuevas reglas.')
table(['Orden propuesto','Entrega y criterio'],[
 ['1 Combate y playtest','Probar contraataque tras parry, espada, arco y alternancia; compilar una build actual y verificar el recorrido.'],
 ['2 Progresión y guardado','Nivel y atributos del personaje; maestría con daño o velocidad; guardar elecciones sin duplicar recompensas.'],
 ['3 Drops y equipamiento','Tiers, modificadores e inventario; comprobar efectos de ambas armas y cambios de vida máxima.'],
 ['4 Heridas y curación','Mostrar máximo curable, cerrar recuperación de heridas y probar una fuente de curación acotada.'],
 ['5 Regiones persistentes','Identificar regiones, fijar mínimos una vez y comprobar carga de chunks, regreso y escalado posterior.'],
 ['6 Apariencias','Destrucción para transmog y colección persistente; integrar skins pagas cuando se defina la plataforma.']],[1.65,5.0])
h('Decisiones todavía abiertas')
p('Atributos definitivos y fórmulas; límites y experiencia de maestría; tiers y tablas de drop; mejora con materiales; porcentaje y recuperación de heridas; reglas exactas del parry; incrementos y disparador de descubrimiento regional; compatibilidad de apariencias y catálogo comercial.')
h('Validación y procedencia')
p('La v0.4 reporta 49 comprobaciones automatizadas aprobadas para su último ajuste de combate. Ese antecedente no valida los sistemas nuevos de la v0.5 ni certifica diversión, balance o rendimiento. La distribución existente señalada en la v0.4 es anterior a sus cambios finales y requiere una build actualizada.')
p('Fuentes: Mismo GDD v0.4 consolidado y decisiones posteriores del autor sobre atributos, tiers, maestría, alcance, equipamiento, parry, heridas, regiones y cosméticos. En caso de conflicto de diseño, esta versión sustituye las propuestas anteriores en esos temas.')

out=ROOT/'Docs'/'Mismo_GDD_v0.5.docx'
D.save(out)
print(out)
