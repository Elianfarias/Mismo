from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

SOURCE = Path(r'C:\Users\elian\Downloads\Mismo_GDD_v0.8.docx')
OUT = Path(r'C:\Users\elian\Mismo\output\gdd\Mismo_GDD_v0.9.docx')
doc = Document(SOURCE)
original = list(doc.element.body.xpath('.//w:p'))
before = [Paragraph(p, doc).text for p in original]
changes = []

def replace(i, text):
    p = Paragraph(original[i], doc)
    changes.append({'paragraph': i, 'before': p.text, 'after': text})
    p.text = text
    return p

def after(i, text, style=None):
    element = OxmlElement('w:p')
    original[i].addnext(element)
    p = Paragraph(element, doc)
    if style:
        p.style = style
    p.text = text
    return p

replacements = {
0: 'Mismo Game Design Document',
1: 'Versión 0.9 consolidada · 3 de octubre de 2026',
2: 'La primera versión completa de Mismo será una aventura individual con cuatro biomas, progresión por familias de armas y un boss por bioma. Este documento define el circuito de exploración, combate y preparación, registra los sistemas actuales y ordena las decisiones y tareas necesarias para completar esa versión. El online, incluido el PvP 3 vs. 3, queda para una etapa posterior.',
3: 'Mismo es un juego de fantasía voxel en tercera persona y sin clases fijas. Las armas determinan las habilidades disponibles. El jugador explora, reúne materiales, fabrica mejoras y consumibles, progresa con sus familias de armas y se prepara para invocar y derrotar al boss de cada bioma. Hay cuatro tipos de bioma: bosque, desierto, nieve y montañas; su contenido y las mazmorras incorporan generación procedural.',
4: 'Se distinguen sistemas actuales, trabajo pendiente, decisiones de diseño abiertas y referencias previas que requieren validación. Hay skins, enemigos, mazmorras, monturas, bestiario, teletransporte desde el mapa y tutorial. Las misiones introductorias y el recorrido de la primera región siguen en construcción. Todavía no hay bosses implementados; el del bosque está en diseño. Los valores numéricos de combate se mantienen como referencias de prueba.',
7: 'Estado y trabajo pendiente',
11: 'Equipamiento por mano y repertorios por familia. El detalle de cuatro estilos, seis habilidades y tres espacios se conserva como referencia a validar. Falta pulir armas y animaciones.',
13: 'Cuatro tipos de bioma definidos, con contenido, mazmorras y distribución de enemigos procedurales. Primera región en construcción. Geografía fija o variable y alcance de cada bioma por cerrar.',
15: 'Menú, audio, minimapa y feedback existentes. Skins y tutorial presentes. Pendientes principales: animaciones de armas, lectura de ataques enemigos, iconos y sonidos.',
17: 'Experiencia por contribución de daño al derrotar monstruos, persistente por familia. Puntos de habilidad cada tres niveles de arma y mejoras de habilidades mediante uso. Inventario y guardado conservan su base previa.',
19: 'Hay una build pública para Windows en itch.io. Cada entrega debe indicar versión y contenido y comprobarse después de los cambios; los paquetes históricos no describen el estado actual.',
23: 'Recolección y fabricación de mejoras, bufos, pociones y un consumible para curarse fuera de combate. Catálogo, costos y balance por completar.',
25: 'Enemigos, bestiario y monturas presentes. Jabalíes y arañas necesitan mejor anticipación y recuperación de ataques. Navegación y poses requieren validación en terreno real.',
28: 'Las reglas vigentes de alcance, progresión y mundo sustituyen las referencias anteriores que entren en conflicto. Los repertorios, números y detalles técnicos conservados ayudan a implementar y probar; no fijan por sí mismos el alcance final. Las decisiones abiertas y los criterios de cierre se reúnen en las secciones 18 y 20.',
33: 'La exploración debe despertar curiosidad por el paisaje y recompensar desvíos. El tutorial explica lo básico y las misiones introductorias guían el aprendizaje de recolección, fabricación y progresión. Cube World clásico orienta cámara, volúmenes y sensación de exploración; Dark Souls, Battlerite, World of Warcraft, Pokémon y Alundra son referencias conceptuales, sin adoptar sus sistemas completos.',
34: 'Circuito de aventura y progresión',
35: 'Comenzar y aprender lo básico → explorar el bioma y sus mazmorras → recolectar y combatir → fabricar, mejorar y elegir habilidades → reunir las condiciones para invocar al boss → vencerlo y avanzar al siguiente bioma. Pueblos, almacenamiento y preparación sostienen las salidas. El orden entre desierto, nieve y montañas queda por confirmar.',
36: 'Cada bioma tendrá un boss que debe invocarse y derrotarse para avanzar. Aún no hay bosses jugables. Se conserva como regla previa por validar que la muerte devuelve al jugador al punto de recuperación sin eliminar materiales, equipo ni progreso persistido. Nodos, botín y recompensas necesitan identidad estable para que cargar o revisitar una zona no duplique premios.',
38: 'El objetivo de la primera versión completa es la aventura individual de los cuatro biomas. La prioridad de producción es cerrar una primera región jugable, desde el tutorial y las misiones iniciales hasta el boss del bosque y su recompensa. El catálogo final de armas, enemigos, mazmorras y recursos de cada bioma debe definirse antes de multiplicar contenido.',
40: 'El online y el PvP 3 vs. 3 se posponen. El cooperativo con anfitrión también queda fuera de la primera versión individual y necesitará un alcance propio. Comercio, refinado avanzado, heridas, transmog, tienda de skins y colección de varias monturas se conservan como propuestas para revisar, sin incorporarlas automáticamente a esta entrega.',
41: 'Se mantiene una aventura sin clases tradicionales. Mundo infinito, terreno destructible, cuevas volumétricas excavables, economía compleja y un catálogo amplio de armas no forman parte del alcance confirmado. Las mazmorras procedurales sí están presentes; no requieren asumir terreno excavable. Las skins actuales se documentan por separado de cualquier futura venta de cosméticos.',
73: 'El modelo se inclina suavemente 6 grados al avanzar y 12 durante sprint, sin inclinar la colisión ni alterar la autoridad de giro del motor. El dash conserva dirección y respeta paredes. Como criterio de recorrido, el acceso al futuro boss no debe exigir doble salto ni dash salvo que se diseñe y enseñe expresamente esa condición.',
102: 'Posture es una defensa enemiga separada de la vida. Como referencia de balance previa, los goblins comienzan con 70, regenera 9 por segundo después de tres segundos sin daño de Posture y al llegar a cero provoca una vulnerabilidad de dos segundos. Golpes adicionales no prolongan esa ruptura; al terminar se restaura la barra. El valor histórico de 150 para un boss es solo una hipótesis de balance: aún no hay bosses implementados.',
125: 'El catálogo previo contempla seis habilidades por familia, un básico fijo y tres espacios Q, E y R, con selección de activas y pasivas. Se conserva como base de repertorio a validar. La regla vigente de desbloqueo es gastar puntos obtenidos cada tres niveles de arma. Los costos, habilidades disponibles al empezar y selección inicial quedan por fijar; las tablas ya no asignan desbloqueos automáticos en niveles 3, 6 y 9.',
127: 'El básico conserva el combo de tres golpes. Los primeros dos permiten ramificar hacia una habilidad compatible o C desde el 65 % de su etapa y durante la transición correspondiente. El buffer de intención es de 0,14 segundos. El golpe final conserva su recuperación. Las referencias históricas Q, E y R orientan la prueba del kit, pero la asignación inicial y los costos de desbloqueo deben revisarse con la progresión por puntos.',
209: 'Daño antes de carga y modificadores. Un toque del básico completa la tensión mínima; mantener incrementa daño y Posture hasta liberar automáticamente a 1,20 s. Su recuperación es 0,35 s. Disparo potente tiene 45 de daño base a Posture y cuesta 12 stamina más 20 Focus. Retirada cuesta 12 stamina y tiene 0,30 s de recuperación; el disparo y el retroceso esperan el final de su preparación. Lluvia de flechas cuesta 20 stamina. Las teclas de la tabla son referencias históricas de prueba, pendientes de confirmar como asignación inicial.',
279: 'La IA del goblin y su variante observa preparaciones ranged al decidir una acción: prioriza una carga disponible si está en rango y cierra más distancia antes de posicionarse. Conserva aviso y dirección fijada; no recibe daño o velocidad extra ocultos por enfrentar al arco. Jabalíes y arañas necesitan mejorar la lectura de preparación y recuperación mediante animación y sonido. Los futuros bosses deberán definir cómo interactúan con Posture y las aperturas.',
280: 'Cuatro biomas y contenido procedural',
281: 'El mundo contiene cuatro tipos de bioma: bosque, desierto, nieve y montañas. El contenido de cada bioma, las mazmorras y las proporciones de aparición de enemigos incorporan generación procedural. Aún debe precisarse qué ocurre con forma, posición y extensión de los biomas entre partidas. La existencia de cuatro tipos fijos no implica por sí sola un mapa idéntico. La semilla, los sectores y las identidades persistentes conservan su función técnica.',
282: 'La estética voxel no implica terreno excavable. La referencia de terreno utiliza alturas escalonadas y mallas con colisión. Hay mazmorras procedurales, cuyo acceso, reglas de generación, recompensas y reinicio deben documentarse. La generación busca variar recorridos y reducir la memorización de una ruta única; no garantiza impedir que se compartan o extraigan contenidos del juego.',
284: 'Bosses secretos muerte y avance',
285: 'El boss del bosque está en diseño y los cuatro bosses de bioma están pendientes de implementación. Para cada uno se deben definir invocación, encuentro, recompensa y condición de avance. El antiguo secreto de una ruina con restauración del 40 % de vida se conserva solo como referencia de prototipo; su vigencia debe revisarse con el contenido actual y no representa el diseño de las mazmorras.',
286: 'El diseño prevé invocar y derrotar un boss por bioma para continuar. El mecanismo concreto que habilita el siguiente bioma, la posibilidad de repetir bosses y el cierre tras el cuarto quedan abiertos. No hay un boss adicional confirmado. Las reglas de muerte y restauración deben preservar recompensas e identidades ya registradas.',
288: 'El recorrido debe conservar orientación, accesibilidad, separación comprensible de encuentros y salidas seguras. Recursos y decoración no deben ocultar ataques ni bloquear caminos. Falta fijar y medir la duración de la primera región y de la aventura completa; el antiguo objetivo de 8 a 12 minutos correspondía a una región de prueba. Rendimiento y ritmo se verifican en una build actual y hardware conocido.',
297: 'El documento previo registra veinte clips adaptados para espada sola, dos espadas y espada con escudo, con ataques, defensa, avance y lanzamiento. Ese inventario es un antecedente técnico; las animaciones siguen siendo el principal pendiente de producción y acabado. Cada familia debe cerrar sus acciones antes de confirmar qué clips se reutilizan, adaptan o necesitan producirse.',
299: 'Menú skins tutorial y opciones',
300: 'MainMenu permite comenzar, abrir opciones y salir. Existen distintas skins; su catálogo, método de selección, desbloqueo y persistencia deben quedar documentados. Durante la partida, el menú general da acceso a Pertenencias, Personaje, Habilidades y Bestiario. Hay un tutorial para lo básico y misiones que guían recolección, fabricación y aprendizaje; las misiones introductorias siguen en construcción. Abrir paneles gestiona cursor y acciones de gameplay, sin equivaler necesariamente a pausa real.',
302: 'AudioManager, AudioEvents y MusicController conservan la estructura de audio del proyecto. Los sonidos existentes se enrutan a efectos. Faltan sonidos de enemigos y revisar la selección de clips, música y señales de combate; disponer de la estructura no certifica una presentación terminada.',
304: 'Mismo tiene una build pública descargable para Windows en itch.io. Mismo-Playtest-2026-09-07.zip queda como antecedente histórico. Antes de distribuir cada actualización se debe generar y comprobar una build que incluya sus cambios, identificar la versión y describir el contenido disponible.',
332: 'Las definiciones de armas, familias, habilidades y reglas pertenecen a Assets/Data, agrupadas por sistema. El arte y las animaciones pertenecen a Assets/Art. Las referencias se configuran mediante campos serializados y el catálogo de Data/System; las búsquedas globales utilizan Mismo.Core.ProjectAssets. Las antiguas referencias a carpetas Resources quedan retiradas de esta documentación.',
333: '8 Personaje familias de armas y dominio de habilidades',
334: 'La progresión distingue atributos generales del personaje, calidad del ejemplar, nivel de la familia y dominio de cada habilidad. Mejorar el objeto mediante crafting, desbloquear habilidades con puntos y obtener una mejora por uso son mecanismos distintos. Los detalles heredados de atributos y tiers se conservan como referencias por validar.',
345: 'Ganar experiencia y niveles de familia, obtener puntos cada tres niveles y elegir habilidades para desbloquear.',
348: 'El diseño previo permite elegir atributos al subir de nivel, inicialmente vida máxima, ataque y defensa o armadura. Deben confirmarse su vigencia, catálogo, puntos, fórmulas y redistribución. Este sistema se documenta aparte de los puntos de habilidad obtenidos por niveles de familia. Defensa y armadura deben unificarse como término si representan la misma estadística.',
349: 'Experiencia niveles y puntos por familia',
350: 'Las armas suben de nivel al derrotar monstruos. La experiencia atribuida a cada arma es proporcional al daño que hizo a ese monstruo y se acumula en su familia. Cambiar una espada por otra de la misma familia conserva el progreso. El catálogo previo diferencia arco, espada sola, dos espadas y espada con escudo; debe revisarse la lista final de familias y sus identificadores al cerrar el catálogo de lanzamiento.',
351: 'Cada tres niveles de arma se otorgan puntos para desbloquear habilidades de su familia. El jugador decide en qué habilidades invertirlos. Este sistema sustituye el desbloqueo automático por alcanzar niveles 3, 6 y 9. Quedan por documentar la cantidad de puntos otorgada en cada hito, costos, requisitos, habilidades iniciales y posibilidad de redistribución. La antigua elección de daño o velocidad mediante puntos debe revisarse antes de conservarla como una regla adicional.',
352: 'La experiencia, nivel, puntos y desbloqueos pertenecen a la familia, no al ejemplar físico. La selección de habilidades también debe conservarse por familia. La atribución de experiencia necesita reglas explícitas para daño periódico, sobreletalidad, enemigos que se curan y otros casos especiales; no se fijan aquí fórmulas aún no acordadas. Los guardados deben preservar el progreso al sustituir un arma, alternar conjuntos y continuar una partida.',
353: 'Mejoras de habilidad por uso y rol de la combinación',
354: 'Usar repetidamente una habilidad permite desbloquear una mejora para esa habilidad. Falta definir qué acciones cuentan como uso válido, los umbrales, el efecto de cada mejora y si existen alternativas excluyentes. El diseño debe premiar la práctica útil y evitar incentivar repeticiones vacías. Las armas y sus combinaciones siguen determinando acciones, ritmo, alcance y oportunidades; las gemas conservan el estado de posibilidad futura.',
377: 'El balance debe evaluar en conjunto atributos del personaje, calidad del ejemplar, habilidades desbloqueadas y mejoras por uso. La velocidad necesita revisar daño por segundo, exposición y frecuencia de efectos. Cualquier bonificación directa de daño o velocidad por maestría queda sujeta a la revisión indicada en la sección 8.',
378: '10 Skins actuales y propuestas de personalización',
379: 'Hay distintas skins en el juego. Falta consolidar su catálogo, selección, condiciones de acceso y persistencia, y dejar explícito si todas son exclusivamente cosméticas. El transmog de armas y las skins pagas que figuran a continuación son propuestas previas por revisar; no se consideran implementados ni comprometidos para la primera versión.',
380: 'Propuesta previa de apariencias mediante armas',
381: 'La propuesta previa consiste en entregar y destruir un ejemplar para aprender permanentemente su apariencia, sin recibir materiales ni oro. Su adopción debe confirmarse junto con las alternativas de conservar, reciclar o vender objetos; estas últimas no se dan por implementadas.',
383: 'Propuesta previa de skins pagas',
384: 'La propuesta anterior contemplaba compras de apariencias cosméticas con dinero real, sin aumentar tier, atributos, maestría ni alcance. No se confirma una tienda para la primera versión. Debe resolverse la vigencia de esta propuesta antes de diseñar su catálogo o integraciones.',
388: 'Los costos de aplicación y desbloqueo de apariencias siguen abiertos. La destrucción del ejemplar y las compras con dinero real pertenecen a las propuestas anteriores y requieren confirmación antes de integrarse al alcance vigente.',
389: '11 Contraataque y propuestas de curación y heridas',
391: 'Propuesta previa por revisar: después de aplicar un parry exitoso, permitir atacar y cortar la recuperación defensiva. Un intento fallido conservaría su recuperación. Su necesidad debe contrastarse con el combate actual antes de declararla parte del cierre.',
394: 'Propuesta previa por revisar, pendiente de implementación: acumular una fracción del daño recibido como herida que la curación normal no recupera. El 10 % era un ejemplo de prueba. Su vigencia debe decidirse junto con pociones, bufos, curación fuera de combate y retorno al pueblo.',
409: '12 Progresión de biomas y escalado enemigo',
410: 'La regla vigente de avance es invocar y derrotar al boss de cada bioma. La cantidad y proporción de enemigos también varían mediante generación procedural. Esto debe distinguirse de su nivel y estadísticas. El antiguo escalado por nivel del personaje se conserva como propuesta para revisar, porque no debe anular la preparación ni la sensación de progreso entre biomas.',
411: 'Referencia previa de mínimo regional',
412: 'Fórmula previa por revisar: mínimo regional = nivel del personaje al descubrir la región + incremento regional.',
414: 'Referencia previa de escalado posterior',
415: 'Fórmula previa por revisar: nivel de los monstruos = máximo entre el mínimo regional guardado y el nivel actual del personaje. Su uso no queda confirmado por el nuevo avance mediante bosses. La maestría de armas y los tiers requieren balance propio.',
425: 'Debe definirse cómo se identifica cada bioma o región, cuándo se considera descubierto y cómo afecta el avance tras un boss a sus encuentros. La precarga técnica de sectores no debe decidir accidentalmente la progresión. La persistencia tiene que conservar descubrimientos, requisitos de invocación, bosses vencidos y accesos habilitados cuando esos sistemas se implementen.',
426: 'Faltan orden y barreras de avance, niveles o bandas de dificultad, pesos de aparición por bioma y reglas de actualización de enemigos. El escalado cooperativo queda fuera de esta versión individual. Las fórmulas y ejemplos previos no sustituyen estas decisiones.',
427: 'El balance debe permitir reconocer la mejora del personaje y preparar al jugador para el siguiente boss. La generación de encuentros necesita límites de densidad, combinaciones comprensibles y acceso fiable a recursos esenciales. Su variación no debe bloquear una partida por falta de un ingrediente o condición de avance.',
439: 'Habilidades muestra repertorios por familia, las opciones equipadas y los requisitos de desbloqueo. La regla actual utiliza puntos obtenidos cada tres niveles de arma; las tarjetas deben indicar costos y condiciones vigentes. Se conserva como referencia la selección de tres espacios Q/E/R y el intercambio de tarjetas para evitar duplicados. Falta validar la presentación de puntos, progreso de uso y mejoras de cada habilidad.',
483: 'La distribución debe ser coherente con cada bioma y respetar terreno, caminos, pueblos, mazmorras y encuentros. La semilla y las identidades conservan los estados persistentes. Hay que garantizar recursos iniciales y materiales necesarios para fabricar, preparar e invocar bosses; generar una posición geométricamente válida no demuestra que el jugador pueda llegar.',
485: 'Existe fabricación de mejoras, bufos, pociones y un consumible para curarse fuera de combate. El banco y las recetas de superficie conservan su función de preparación. Deben documentarse recetas, cantidades, efectos, restricciones de uso y fuentes de ingredientes de cada bioma. Las recetas de la tabla siguiente son referencias iniciales del sistema, no el catálogo completo actual.',
497: 'Está confirmado un consumible de curación fuera de combate. El ungüento conserva como referencia previa 40 de vida durante 5 segundos, sin superar el máximo, sin consumo con vida llena o efecto idéntico activo y con interrupción al recibir daño. Deben comprobarse estos valores y distinguir sus reglas de las pociones y bufos. Las heridas continúan como propuesta por revisar.',
499: 'Mazmorras recursos y ampliación de recetas',
500: 'Las mazmorras procedurales forman parte del juego actual. Falta consolidar su propósito, acceso, estructura, variedad, recompensas y reglas de persistencia o regeneración. También debe definirse su relación con los materiales y requisitos de invocación de bosses. El diseño necesita recorridos alcanzables y una salida segura; la estética voxel no exige excavación volumétrica.',
501: 'El escudo básico T1 figuraba como fabricable mediante REC-09, con 2 maderas y 4 hierros; sus costos y el resto del catálogo deben revisarse. Las combinaciones de armas resultan del equipamiento de ejemplares individuales. El catálogo actual de mejoras, bufos y pociones requiere consolidación sin sustituir todas las recompensas propias de bosses, mazmorras y secretos.',
504: 'Monturas y bestiario están presentes. La regla previa de domesticación usa una probabilidad configurable al vencer a una criatura elegible. Deben confirmarse su vigencia, especies habilitadas y condiciones actuales. Se conserva la distinción entre contar victorias y cumplir una condición de desbloqueo.',
509: 'El catálogo previo incluía jabalí, araña, goblin, gólem y guardián, con el jabalí como primera montura. Hay distintas especies y monturas en el estado actual; falta consolidar el catálogo real y las especies habilitadas. Una ficha de guardián no demuestra que exista un boss jugable. Las fichas deben revelarse al descubrir la criatura según la regla de exploración vigente.',
517: 'El taller de monturas permite previsualizar personaje y criatura en una escena aislada y ajustar asiento, escala, clip y huesos. La pose inicial de jabalí requiere revisión en movimiento y pendientes. El catálogo actual de monturas debe indicar qué poses están resueltas y cuáles faltan; conservar un taller no equivale a tener todas las animaciones finales.',
526: 'Assets/Data, agrupado por sistema; definiciones de objetos, armas, tamaños, apilado e iconos.',
528: 'Assets/Data, sistema de recolección; nodos, tiempos, recompensas y distribución.',
530: 'Assets/Data, sistema de recetas; costos, productos y mejoras.',
534: 'Assets/Data, sistema de criaturas y bestiario; especies, elegibilidad y páginas.',
538: 'Assets/Data, sistema de idiomas; tablas de español e inglés y claves estables.',
539: 'Las definiciones compartidas contienen configuración; el guardado conserva objetos, familias de armas, puntos y desbloqueos, progreso de uso de habilidades, selección, cofre, nodos, descubrimientos y monturas. Debe documentarse y validarse además la persistencia de skins, tutorial, misiones, pueblos de viaje, mazmorras y futuros bosses. Las migraciones deben preservar partidas compatibles.',
541: '18 Prioridades de producción y criterios de cierre',
542: 'La prioridad es definir y completar una primera región individual representativa. Faltan todos los bosses y el del bosque está en diseño. El principal bloqueo de producción son las animaciones, junto con el acabado de combate, iconos y sonido. Las decisiones pendientes deben resolverse antes de ampliar el repertorio o prometer la aventura completa.',
548: 'Fijar objetivo, tutorial, misiones, recursos, mazmorra, preparación e invocación del boss del bosque, con inicio y cierre comprensibles.',
547: 'Definir el recorrido de la primera región',
550: 'Diseñar e implementar el boss del bosque',
551: 'Cerrar invocación, ataques, respuestas del jugador, recompensa y condición de avance; después verificar el encuentro completo.',
553: 'Cerrar armas y enemigos de la región',
554: 'Definir el repertorio necesario y verificar diferencias de juego. Jabalíes y arañas deben anunciar ataques y ofrecer recuperaciones legibles.',
556: 'Completar animaciones iconos y sonido',
557: 'Revisar anticipación, impacto, recuperación, agarres y transiciones; asignar iconos y sonidos que comuniquen acciones y estados.',
559: 'Validar el recorrido en una build',
560: 'Un jugador nuevo aprende, recolecta, fabrica, progresa, guarda, continúa y supera la región sin asistencia del desarrollador.',
562: 'Completar los otros tres biomas y bosses',
563: 'Cada bioma tiene identidad, recursos, encuentros, contenido procedural y boss propios; los cuatro forman una progresión completa.',
565: 'Alcance y calidad de la entrega',
566: 'Comprobar familias, puntos, mejoras por uso, viaje, monturas, misiones y persistencia; medir rendimiento y comunicar contenido y versión.',
574: 'Verificar que cada hito de tres niveles de arma otorga los puntos definidos, que desbloquear una habilidad descuenta el costo correcto y que la experiencia por contribución de daño se conserva por familia. Validar también uso, mejora, selección y guardado sin repetir premios al cargar.',
576: 'Antecedentes de validación',
577: 'El GDD anterior registra 105 comprobaciones de inventario y guardado, 15 de desbloqueo y selección y 150 de importación de assets, además de compilación y revisión de veinte clips. Son antecedentes de la versión anterior y no certifican la nueva progresión por puntos, las mejoras por uso, las mazmorras o el avance mediante bosses.',
578: 'También se conservan los antecedentes de recolección y criaturas de las versiones 0.6 y 0.7. La versión 0.9 reemplaza las reglas anteriores de desbloqueo y alcance que contradigan sus secciones vigentes. El cierre requiere una build actual y recorridos reales; los números históricos no se presentan como pruebas ejecutadas para esta actualización.',
580: 'Diseño previo en borrador, pendiente de implementación y de confirmar dentro del catálogo de la primera versión. Hacha sola y Dos hachas se conservarían como familias distintas con identificadores propios. Si se incorporan, deben aplicar la progresión vigente por contribución de daño, puntos cada tres niveles y mejoras por uso. Las habilidades siguientes son propuestas de repertorio; no fijan costos ni niveles automáticos de desbloqueo.',
}
for i, text in replacements.items():
    replace(i, text)

# Retire old automatic mastery thresholds from all six repertoire tables.
for table_index in (4, 6, 7, 8, 19, 20):
    table = doc.tables[table_index]
    table.cell(0, 2).text = 'Desbloqueo'
    for row in table.rows[1:]:
        row.cells[2].text = 'Costo por definir'

# Clarify the inherited regional scaling example.
doc.tables[13].cell(0, 0).text = 'Ejemplo previo con incremento de 15'
doc.tables[13].cell(0, 1).text = 'Nivel resultante por revisar'

# Complete the current-state matrix without treating missing details as bugs.
for cells in [
    ('Bosses y avance', 'Un boss previsto por bioma. Ninguno implementado; bosque en diseño. Invocación, encuentros, recompensas y cierre pendientes.'),
    ('Mapa y viaje', 'Teletransporte al pulsar un pueblo en el mapa. Descubrimiento, costos, restricciones y relación con monturas por documentar.'),
    ('Aprendizaje y misiones', 'Tutorial básico presente. Misiones de recolección y crafting en construcción; falta cerrar su conexión con el avance de la primera región.'),
]:
    row = doc.tables[0].add_row()
    for cell, text in zip(row.cells, cells):
        cell.text = text

# Add short explanations at the point of use rather than an amendment-only appendix.
after(296, 'Antes de producir o buscar clips, cada acción debe tener una ficha con función, preparación visible, dirección y alcance, momento de impacto, respuesta permitida y recuperación. Ese repertorio determina la lista de animaciones necesarias. La sección 20 reúne los campos a cerrar.')
after(295, 'El mapa permite teletransportarse al pulsar un pueblo. Falta documentar si exige descubrirlo antes, si tiene costo, si puede usarse durante combate o dentro de una mazmorra y qué utilidad conserva la montura en los trayectos locales. No se asumen restricciones aún no acordadas.')
after(354, 'Las mejoras por uso deben registrarse por habilidad y conservarse cuando se cambia de ejemplar de la misma familia. Falta acordar si se activan automáticamente, requieren selección o consumen algún recurso. Esas decisiones son independientes del gasto de puntos para desbloquear la habilidad.')
after(515, 'El límite de una montura, ausencia de salto y combate montado y reglas de seguir o esperar son referencias previas que requieren comprobarse frente al catálogo actual. No se amplían ni eliminan automáticamente por existir nuevas monturas.')

# Add domain-specific closure rules; these are design criteria, not test results.
after(575, 'Validar que tutorial y misiones registren el progreso una sola vez y se conserven al continuar. Comprobar que teletransporte, monturas y mazmorras no omitan requisitos de avance. Cuando existan bosses, guardar y cargar debe conservar invocación, victoria, recompensa y acceso habilitado sin duplicaciones.', 'List Bullet')

doc.add_heading('20 Decisiones para cerrar la primera versión', 1)
doc.add_paragraph('La aventura individual de cuatro biomas es el alcance confirmado. Esta lista permite completar sus reglas antes de estimar producción, animaciones o fechas. Una decisión pendiente no equivale a una funcionalidad ausente; los estados de implementación se indican cuando están confirmados.')

doc.add_heading('Biomas y bosses', 2)
doc.add_paragraph('Bosque, desierto, nieve y montañas son los cuatro tipos de bioma previstos. El bosque concentra el trabajo inicial. Deben definirse el orden restante, la parte fija o variable del mapa, la identidad visual y jugable de cada bioma y sus recursos y enemigos propios.')
table = doc.add_table(rows=1, cols=3)
for c, t in zip(table.rows[0].cells, ['Bioma', 'Boss al 3 de octubre', 'Definiciones pendientes']): c.text=t
for data in [
    ('Bosque', 'En diseño; no implementado', 'Invocación, aprendizaje que evalúa, ataques, recompensa y avance.'),
    ('Desierto', 'No implementado', 'Concepto, invocación, encuentro, recompensa y lugar en la progresión.'),
    ('Nieve', 'No implementado', 'Concepto, invocación, encuentro, recompensa y lugar en la progresión.'),
    ('Montañas', 'No implementado', 'Concepto, invocación, encuentro, recompensa y lugar en la progresión.'),
]:
    for c,t in zip(table.add_row().cells,data): c.text=t
doc.add_paragraph('El plan incluye un boss por bioma. Falta definir el cierre tras vencer al cuarto y qué actividades continúan disponibles. No hay un boss adicional confirmado. La implementación del resto del contenido de cada bioma debe registrarse por separado del estado de su boss.')

doc.add_heading('Armas habilidades y recompensas', 2)
doc.add_paragraph('Cerrar el catálogo de familias y las habilidades de cada una. Para los puntos obtenidos cada tres niveles, fijar cantidad por hito, costos, requisitos, habilidades iniciales y redistribución. Para las mejoras por uso, definir uso válido, umbrales, efectos y forma de activación. Mantener separadas la calidad del objeto y la progresión persistente de familia.')
doc.add_paragraph('Cada recurso, receta, drop y recompensa necesita un propósito en la preparación o exploración. Documentar costos y duración de mejoras, bufos y pociones, apilado de efectos y restricciones de uso. El consumible de curación fuera de combate debe tener una regla propia y comprensible.')

doc.add_heading('Mazmorras encuentros y desplazamiento', 2)
doc.add_paragraph('Definir objetivo, recompensas, acceso y cierre de las mazmorras; decidir cuándo se generan o reinician y qué permanece al guardar. La variación procedural debe mantener rutas y recompensas alcanzables. Para enemigos, fijar pesos, límites y combinaciones por bioma, evitando confundir frecuencia de aparición con escalado de nivel.')
doc.add_paragraph('Precisar las condiciones del teletransporte desde pueblos del mapa y su relación con misiones, combate, mazmorras y progresión por bosses. Completar el catálogo de monturas, forma de obtención, funciones y reglas de recuperación ante atascos. Confirmar también selección y persistencia de skins.')

doc.add_heading('Primera región y aprendizaje', 2)
doc.add_paragraph('Definir una secuencia que enseñe movimiento y combate, recolección, fabricación, elección de habilidades y preparación para el boss del bosque. Las misiones deben orientar con objetivos y recompensas comprensibles. Cerrar el objetivo de la región y medir su duración en pruebas antes de prometer tiempos de juego.')
doc.add_paragraph('Criterio de cierre: un jugador nuevo puede iniciar, comprender sus próximos objetivos, reunir materiales, fabricar, ganar progreso de arma, preparar la invocación y vencer al boss sin asistencia del desarrollador. Guardar y continuar conserva cada paso confirmado. El recorrido debe probarse con distintas variaciones del contenido procedural.')

doc.add_heading('Repertorio y producción de animaciones', 2)
doc.add_paragraph('La principal dificultad actual es producir o conseguir animaciones. Primero se define el repertorio necesario de armas, enemigos y bosses; luego se registra qué clips existen, cuáles se pueden adaptar y cuáles faltan. Jabalíes y arañas tienen prioridad por la poca claridad actual de sus ataques.')
for text in [
    'Función de la acción: qué decisión permite al jugador o qué situación crea el enemigo.',
    'Preparación: pose, movimiento y señal sonora que anuncian intención, dirección y alcance.',
    'Impacto: momento del contacto, zona afectada y coincidencia entre animación, daño y efectos.',
    'Respuesta: posibilidad de esquivar, bloquear, interrumpir o reposicionarse, según el ataque.',
    'Recuperación y transiciones: cuánto se expone el atacante y cómo vuelve al movimiento o encadena otra acción.',
]: doc.add_paragraph(text, style='List Bullet')
doc.add_paragraph('La revisión debe hacerse en combate real: agarres, pies, orientación, velocidad, interrupciones y lectura de varios enemigos. Íconos y sonidos acompañan el mismo repertorio. No se fijan cantidades de animaciones ni encargos hasta cerrar las acciones necesarias.')

doc.add_heading('Propuestas que requieren confirmación', 2)
doc.add_paragraph('Heridas, comercio, refinado avanzado, transmog de armas, skins pagas, gemas, colección de varias monturas y el repertorio de hachas conservan su historial de diseño. Hay que decidir cuáles siguen vigentes y cuáles se posponen. Online, cooperativo y PvP 3 vs. 3 quedan para después de la aventura individual.')

# Preserve the original Calibri design, regularize black headings, and make tables readable.
for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
    if name in doc.styles:
        style = doc.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style.element.rPr is not None:
            for color in style.element.rPr.findall(qn('w:color')):
                for attr in ('themeColor','themeTint','themeShade'):
                    color.attrib.pop(qn('w:'+attr),None)

for p in doc.paragraphs:
    p.paragraph_format.page_break_before = False
    for page_break in list(p._p.xpath('.//w:br[@w:type="page"]')):
        page_break.getparent().remove(page_break)
    if p.style.name.startswith('Heading') or p.style.name == 'Title':
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        ppr = p._p.pPr
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
        for run in p.runs:
            run.font.color.rgb=RGBColor(0,0,0)
            run.font.underline=False
    p.paragraph_format.widow_control=True
    if p.text == '20 Decisiones para cerrar la primera versión':
        p.paragraph_format.page_break_before = True

doc.styles['Heading 1'].paragraph_format.page_break_before=False
doc.styles['Heading 1'].paragraph_format.space_before=Pt(20)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(6)

width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
for ti, table in enumerate(doc.tables):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    n=len(table.columns)
    if ti==0: ratios=[0.26,0.74]
    elif ti==18: ratios=[0.12,0.30,0.58]
    elif ti==21: ratios=[0.17,0.25,0.58]
    elif ti in (4,6,7,8,19,20): ratios=[0.21,0.13,0.19,0.47]
    elif n==2: ratios=[0.31,0.69]
    elif n==3: ratios=[0.26,0.34,0.40]
    else: ratios=[1/n]*n
    for col, ratio in zip(table.columns,ratios): col.width=int(width*ratio)
    props=table._tbl.tblPr
    borders=props.find(qn('w:tblBorders'))
    if borders is None: borders=OxmlElement('w:tblBorders');props.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=borders.find(qn('w:'+edge))
        if e is None:e=OxmlElement('w:'+edge);borders.append(e)
        e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9')
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr()
        if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
        if ri==0 and trpr.find(qn('w:tblHeader')) is None:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=int(width*ratios[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            cell_borders=tcpr.find(qn('w:tcBorders'))
            if cell_borders is not None:tcpr.remove(cell_borders)
            cell_borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                border=OxmlElement('w:'+edge)
                border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9')
                cell_borders.append(border)
            tcpr.append(cell_borders)
            shd=tcpr.find(qn('w:shd'))
            if shd is None:shd=OxmlElement('w:shd');tcpr.append(shd)
            shd.set(qn('w:fill'),'E7EDF2' if ri==0 else 'FFFFFF')
            margins=tcpr.find(qn('w:tcMar'))
            if margins is None:margins=OxmlElement('w:tcMar');tcpr.append(margins)
            for side,value in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
                item=margins.find(qn('w:'+side))
                if item is None:item=OxmlElement('w:'+side);margins.append(item)
                item.set(qn('w:w'),value);item.set(qn('w:type'),'dxa')
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=(ri==0 or (len(table.rows)<=5 and ri<len(table.rows)-1))
                p.paragraph_format.keep_together=True
                for run in p.runs:
                    run.font.name='Calibri';run.font.size=Pt(10)
                    run.font.color.rgb=RGBColor(0,0,0)
                    if ri==0:run.font.bold=True
    next_element=table._tbl.getnext()
    if next_element is not None and next_element.tag==qn('w:p'):
        following=Paragraph(next_element,doc)
        if not following.style.name.startswith('Heading'):
            following.paragraph_format.space_before=Pt(7)

doc.core_properties.title='Mismo Game Design Document'
doc.core_properties.subject='Aventura individual y reglas de progresión de la versión 0.9'
doc.core_properties.version='0.9'
doc.core_properties.modified=datetime(2026,10,3,12,0,0,tzinfo=timezone.utc)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
(OUT.parent/'qa_changes.json').write_text(json.dumps({'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'output':str(OUT),'changes':changes},ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Saved: {OUT}')
print(f'Updated paragraphs: {len(changes)}; tables: {len(doc.tables)}')

# Focused content checks target the actual contradictions addressed in this revision.
full='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for required in ['Cada tres niveles de arma','proporcional al daño','No hay un boss adicional confirmado','No implementado','teletransportarse','mazmorras procedurales','Versión 0.9']:
    assert required in full, required
for retired in ['El boss conserva sus patrones','No existe dungeon volumétrico','se desbloquean en niveles 3, 6 y 9','Assets/Resources/','Versión 0.7 consolidada']:
    assert retired not in full, retired
print('Content checks passed')
